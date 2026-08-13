import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from marine_uw_agent.history import filter_same_vessel_quotes
from marine_uw_agent.models import QuotationSlip, ReviewState, RunInputs
from marine_uw_agent.normalizers import classify_match, normalize_company_name, normalize_imo, normalize_vessel_name
from marine_uw_agent.paths import create_run_paths, evidence_path
from marine_uw_agent.quotation_history_adapter import QuotationHistoryAdapter
from marine_uw_agent.quotation_extractor import extract_quotation_from_text
from marine_uw_agent.restrictions import cargo_restriction_review
from marine_uw_agent.run import run_workflow
from marine_uw_agent.sanctions import LocalSanctionsIndex
from marine_uw_agent.workbook import WorkbookPopulationError, populate_workbook
from scripts.run_uw_web_evidence import eu_tracker_terms_from_manifest, extract_equasis_ship_fields, management_rows_from_dom


class MarineUWAgentTests(unittest.TestCase):
    def test_eu_tracker_uses_all_equasis_management_companies(self):
        args = SimpleNamespace(
            imo="9342865",
            vessel_name="Fallback Vessel",
            commercial_manager="Fallback Manager",
            registered_owner="Fallback Owner",
        )
        manifest = {
            "review": {
                "imo": "9342865",
                "vesselName": "TRAVERSE SINGAPORE",
                "commercialManager": "Goldenking Ship Management",
                "registeredOwner": "Traverse Shipping Co Ltd",
            },
            "equasis": {"details": {"commercialManager": "Goldenking Ship Management"}},
            "fleet": {
                "manager": [
                    {
                        "role": "Ship manager",
                        "company": "Second Management Company",
                        "company_imo": "7654321",
                        "address": "1 Harbour Road, Singapore",
                    }
                ],
                "owner": [
                    {
                        "role": "Registered owner",
                        "company": "Traverse Shipping Co Ltd",
                        "company_imo": "1234567",
                    }
                ],
            },
        }

        terms = eu_tracker_terms_from_manifest(manifest, args)
        values = [term.value for term in terms]

        self.assertIn("TRAVERSE SINGAPORE", values)
        self.assertIn("Goldenking Ship Management", values)
        self.assertIn("Second Management Company", values)
        self.assertNotIn("7654321", values)
        self.assertNotIn("1 Harbour Road, Singapore", values)
        self.assertEqual(values.count("Traverse Shipping Co Ltd"), 1)

    def test_sanctions_terms_keep_same_company_for_distinct_equasis_roles(self):
        args = SimpleNamespace(imo="9342865", vessel_name="", commercial_manager="", registered_owner="")
        manifest = {
            "review": {"imo": "9342865", "vesselName": "TRAVERSE SINGAPORE"},
            "equasis": {"management": [
                {"role": "Ship manager/Commercial manager", "company": "GOLDENKING SHIP MANAGEMENT"},
                {"role": "ISM Manager", "company": "GOLDENKING SHIP MANAGEMENT"},
                {"role": "Registered owner", "company": "TRAVERSE SHIPPING CO LTD-MAI"},
            ]},
        }
        terms = eu_tracker_terms_from_manifest(manifest, args)
        goldenking = [term for term in terms if term.value == "GOLDENKING SHIP MANAGEMENT"]
        self.assertEqual([term.label for term in goldenking], ["Ship manager/Commercial manager", "ISM Manager"])

    def test_equasis_dom_management_table_keeps_roles_and_fields(self):
        args = SimpleNamespace(imo="9342865", vessel_name="Fallback", flag="Panama")
        dom = {"tables": [{
            "headers": ["IMO number", "Role", "Name of company", "Address", "Date of effect", "Details"],
            "rows": [
                ["6212304", "Ship manager/Commercial manager", "GOLDENKING SHIP MANAGEMENT", "Guangzhou, China", "since 16/10/2025", ""],
                ["6212304", "ISM Manager", "GOLDENKING SHIP MANAGEMENT", "Guangzhou, China", "since 16/10/2025", ""],
                ["0358925", "Registered owner", "TRAVERSE SHIPPING CO LTD-MAI", "Guangzhou, China", "since 16/10/2025", ""],
            ],
        }]}
        rows = management_rows_from_dom(dom, {"vesselName": "TRAVERSE SINGAPORE", "flag": "Panama"}, args)
        self.assertEqual([row["role"] for row in rows], ["Ship manager/Commercial manager", "ISM Manager", "Registered owner"])
        self.assertEqual(rows[2]["company_imo"], "0358925")
        self.assertEqual(rows[0]["address"], "Guangzhou, China")

    def test_equasis_ship_particular_fields_are_scraped(self):
        raw = """Flag\n(Panama)\nCall Sign\n3EDM7\nMMSI\n371626000\nGross tonnage\n39738
        (since 01/12/2005)\nDWT\n76619\nType of ship\nBulk Carrier\nYear of build\n2005
        Status\nIn Service/Commission\nLast update of ship particulars\n2026-08-11"""
        fields = extract_equasis_ship_fields(raw)
        self.assertEqual(fields["Flag"], "Panama")
        self.assertEqual(fields["MMSI"], "371626000")
        self.assertEqual(fields["Type of ship"], "Bulk Carrier")

    def test_quotation_text_extraction(self):
        text = """
        Quotation Date: 2026-06-17
        Type of Coverage: Marine War Hull
        Assured: Example Assured Ltd
        Vessel Name: DUBAI TOWER
        IMO: 9433066
        Route: Singapore to Fujairah via Gulf of Oman
        Listed Area: JWLA-033 Gulf of Oman
        Net Rate: 0.025%
        Premium: USD 25,000
        Payment Terms: 30 days
        """

        slip = extract_quotation_from_text("ART-1", text)

        self.assertEqual(slip.vessel_name, "DUBAI TOWER")
        self.assertEqual(slip.imo, "9433066")
        self.assertEqual(slip.listed_area, "JWLA-033 Gulf of Oman")

    def test_normalization_helpers(self):
        self.assertEqual(normalize_imo("IMO 9433066"), "9433066")
        self.assertEqual(normalize_vessel_name("DUBAI-TOWER"), "dubaitower")
        self.assertEqual(normalize_company_name("Transfar Shipping Pte Ltd"), "transfar shipping")

    def test_sanctions_match_classification(self):
        self.assertEqual(classify_match("9433066", "Vessel IMO 9433066", exact_identifier=True), "confirmed_match")
        self.assertEqual(classify_match("Blocked UK Owner Ltd", "Blocked UK Owner Limited", score=100), "likely_match")
        self.assertEqual(classify_match("Blocked Owner", "Blocked Owners", score=88), "possible_match")
        self.assertEqual(classify_match("Clear Owner", "Different Company", score=20), "clear")

    def test_prior_quotation_date_window(self):
        current = {"imo": "9433066", "vessel_name": "Dubai Tower"}
        quotes = [
            {"quotation_reference": "Q1", "quotation_date": "2026-06-01", "imo": "9433066"},
            {"quotation_reference": "Q2", "quotation_date": "2026-03-17", "vessel_name": "DUBAI TOWER"},
            {"quotation_reference": "Q3", "quotation_date": "2026-03-16", "imo": "9433066"},
        ]

        matches = filter_same_vessel_quotes(current, quotes, date(2026, 6, 17))

        self.assertEqual([item["quotation_reference"] for item in matches], ["Q1", "Q2"])

    def test_local_sanctions_index(self):
        with tempfile.TemporaryDirectory() as tmp:
            snapshot = Path(tmp) / "snapshot.json"
            snapshot.write_text(
                json.dumps(
                    {
                        "generated_at": "2026-06-01T00:00:00Z",
                        "records": [{"source": "uk", "name": "Blocked UK Owner Limited", "program": "Test", "reference": "UK1"}],
                    }
                ),
                encoding="utf-8",
            )
            index = LocalSanctionsIndex(snapshot)
            from marine_uw_agent.models import Party

            matches = index.search_party(Party("Blocked UK Owner Ltd", "owner"))

            self.assertEqual(matches[0].status, "likely_match")
            self.assertEqual(matches[0].unique_id, "UK1")

    def test_cargo_missing_is_information_insufficient(self):
        results = cargo_restriction_review({}, {"eu": "eu", "uk": "uk", "ofac": "ofac"})
        self.assertTrue(all(item.status == "information_insufficient" for item in results))

    def test_evidence_path_generation(self):
        with tempfile.TemporaryDirectory() as tmp:
            paths = create_run_paths(Path(tmp) / "output", "ART/REF 1")
            path = evidence_path(paths, "UK Sanctions", "DUBAI TOWER")
            self.assertTrue(str(path).endswith("UK_Sanctions_DUBAI_TOWER.png"))

    def test_workbook_population_requires_openpyxl_when_missing(self):
        inputs = RunInputs("ART-1", date(2026, 6, 17), Path("missing.xlsx"), Path("output"))
        state = ReviewState(inputs, "test", "2026-06-17T00:00:00Z", quotation=QuotationSlip("ART-1", vessel_name="DUBAI TOWER", imo="9433066"))
        try:
            import openpyxl  # noqa: F401
        except Exception:
            with self.assertRaises(WorkbookPopulationError):
                populate_workbook(state, Path("output/test.xlsx"))
        else:
            with tempfile.TemporaryDirectory() as tmp:
                output = populate_workbook(state, Path(tmp) / "test.xlsx")
                self.assertTrue(output.exists())

    def test_quotation_history_adapter_normalizes_rows(self):
        adapter = object.__new__(QuotationHistoryAdapter)
        row = adapter._normalize_row(
            {
                "Quotation Ref no.": "ART-2026031605-MW",
                "Broker": "28-BMC Asia/JS Risks",
                "IMO": "9307736",
                "Listed Area": "Persian Gulf/Gulf of Oman",
                "Type of Coverage": "Marine Hull War",
                "Gross Rate": "0.50%",
                "Discount": "12.5%",
                "Net Rate": "0.4375%",
                "Premium": "USD 200,666.01",
                "Bound or not": "Bound",
                "Quotation Date": "2026-03-16",
            }
        )

        self.assertEqual(row["quotation_reference"], "ART-2026031605-MW")
        self.assertEqual(row["vessel_imo"], "9307736")
        self.assertEqual(row["coverage_type"], "Marine Hull War")
        self.assertEqual(row["bound"], "Bound")

    def test_quotation_history_adapter_normalizes_api_rows(self):
        adapter = object.__new__(QuotationHistoryAdapter)
        row = adapter._normalize_row(
            {
                "quotationNo": "ART-2026031605-MW",
                "brokerName": "28-BMC Asia/JS Risks",
                "shipImo": "9307736",
                "classType": "Marine Hull War",
                "grossRate": "0.50%",
                "netRate": "0.4375%",
                "netPremium": "USD 200,666.01",
                "statusName": "Bound",
                "createTime": "2026-03-16",
            }
        )

        self.assertEqual(row["quotation_reference"], "ART-2026031605-MW")
        self.assertEqual(row["broker"], "28-BMC Asia/JS Risks")
        self.assertEqual(row["vessel_imo"], "9307736")
        self.assertEqual(row["net_rate"], "0.4375%")

    def test_aquamarine_template_population_preserves_excel_output(self):
        template = Path("data/raw/UW Review - draft - Aquamarine - v2.xlsx")
        if not template.exists():
            self.skipTest("Aquamarine template not present")
        with tempfile.TemporaryDirectory() as tmp:
            inputs = RunInputs("ART-NEW", date(2026, 7, 17), template, Path(tmp) / "output")
            state = ReviewState(
                inputs,
                "test",
                "2026-07-17T00:00:00Z",
                quotation=QuotationSlip(
                    "ART-NEW",
                    vessel_name="TRAVERSE SINGAPORE",
                    imo="9342865",
                    flag="Panama",
                    vessel_type="Bulk Carrier",
                    listed_area="Gulf of Aden/Red Sea/Indian Ocean",
                ),
                quotation_history=[
                    {
                        "quotation_reference": "ART-OLD",
                        "broker": "Broker A",
                        "vessel_imo": "9342865",
                        "listed_area": "Gulf of Aden",
                        "coverage_type": "Marine Hull War",
                        "gross_rate": "0.10%",
                        "discounts": "0%",
                        "net_rate": "0.10%",
                        "premium": "USD 1,000",
                        "bound": "Not bound",
                        "quotation_date": "2026-06-01",
                    }
                ],
            )
            output = populate_workbook(state, Path(tmp) / "review.xlsx")
            import openpyxl

            wb = openpyxl.load_workbook(output, data_only=True)
            ws = wb["UW Info"]
            self.assertEqual(ws["B1"].value, "ART-NEW")
            self.assertEqual(ws["A13"].value, "Quotation Ref no.")
            self.assertEqual(ws["A14"].value, "ART-OLD")
            self.assertEqual(ws["C14"].value, "9342865")

    def test_offline_workflow_marks_missing_information(self):
        with tempfile.TemporaryDirectory() as tmp:
            quote_text = Path(tmp) / "quote.txt"
            quote_text.write_text("Vessel Name: DUBAI TOWER\nIMO: 9433066\nListed Area: JWLA-033\n", encoding="utf-8")
            inputs = RunInputs("ART-1", date(2026, 6, 17), Path("missing.xlsx"), Path(tmp) / "output")

            state = run_workflow(inputs, quotation_text_path=quote_text, offline=True)

            self.assertIn("missing cargo details", state.open_issues)
            self.assertIn("missing port history", state.open_issues)


if __name__ == "__main__":
    unittest.main()
