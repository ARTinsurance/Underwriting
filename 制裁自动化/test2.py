import time
import os
import pandas as pd
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager
from docx import Document
from docx.shared import Inches
from PIL import Image



# ========== 参数配置 ==========
EXCEL_PATH = '制裁自动化\公司名称.xlsx'    # Excel文件路径
OUTPUT_DIR = 'screenshots'       # 截图保存文件夹
WORD_PATH = 'OFAC_Search.docx'   # Word文件路径

# 截图区域坐标
CROP_X = 500
CROP_Y = 70
CROP_WIDTH = 900
CROP_HEIGHT = 700

# ========== 创建截图文件夹 ==========
os.makedirs(OUTPUT_DIR, exist_ok=True)

# ========== 读取Excel ==========
df = pd.read_excel(EXCEL_PATH)
company_list = df.iloc[:, 0].dropna().tolist()

# ========== 启动浏览器 ==========
options = Options()
options.add_argument('--start-maximized')

# 如果你不想显示浏览器界面，可取消下面这行注释
# options.add_argument('--headless')

service = Service(ChromeDriverManager().install())
driver = webdriver.Chrome(service=service, options=options)

# ========== 定义截图函数 ==========
def search_and_screenshot(company_name, index):
    driver.get("https://sanctionssearch.ofac.treas.gov/")

    # 等待页面加载
    time.sleep(2)

    # 定位输入框并输入公司名
    input_box = driver.find_element(By.ID, 'ctl00_MainContent_txtLastName')
    input_box.clear()
    input_box.send_keys(company_name)

    # 点击Search按钮
    search_btn = driver.find_element(By.ID, 'ctl00_MainContent_btnSearch')
    search_btn.click()

    # 等待搜索结果加载
    time.sleep(3)

    # 将页面缩小到75%
    driver.execute_script("document.body.style.zoom='75%'")
    time.sleep(1)  # 等待缩放渲染

    # 全屏截图路径
    full_screenshot_path = os.path.join(OUTPUT_DIR, f'full_screenshot_{index}.png')
    driver.save_screenshot(full_screenshot_path)

    # 打开全屏图片并裁剪
    with Image.open(full_screenshot_path) as img:
        cropped = img.crop((CROP_X, CROP_Y, CROP_X + CROP_WIDTH, CROP_Y + CROP_HEIGHT))
        cropped_path = os.path.join(OUTPUT_DIR, f'screenshot_{index}.png')
        cropped.save(cropped_path)

    # 删除全屏截图文件（如不需要可保留）
    os.remove(full_screenshot_path)

    return cropped_path

# ========== 搜索并截图 ==========
screenshot_paths = []

for idx, company in enumerate(company_list, start=1):
    print(f"正在查询: {company}")
    try:
        path = search_and_screenshot(company, idx)
        screenshot_paths.append((company, path))
    except Exception as e:
        print(f"查询 {company} 出错: {e}")
        screenshot_paths.append((company, None))

# ========== 写入Word ==========
doc = Document()
doc.add_heading('OFAC Sanctions Search Results', level=1)

for company, img_path in screenshot_paths:
    doc.add_heading(company, level=2)
    if img_path and os.path.exists(img_path):
        doc.add_picture(img_path, width=Inches(6))
    else:
        doc.add_paragraph("No screenshot available.")

# 保存Word文档
doc.save(WORD_PATH)

print("全部完成！")

driver.quit()



