import os
import time
import pandas as pd
from selenium import webdriver
from selenium.webdriver.edge.service import Service
from selenium.webdriver.edge.options import Options
from docx import Document
from PIL import ImageGrab
import pyautogui
import pyperclip
from docx.shared import Inches

# ---------- 参数配置 ----------

# Excel路径
excel_path = '制裁自动化\公司名称.xlsx'

# PDF文件路径(只能绝对路径)
pdf_path = r'D:\桌面\制裁自动化(1)\制裁自动化\20250708-European UnionConsolidated Financial Sanctions List.pdf'

# Word保存路径
word_path = 'search_results.docx'

# 截图区域坐标 (left, top, right, bottom)
# 根据你的屏幕和 Edge 窗口位置自行调整
screenshot_box = (0, 0, 1920, 1010)


# Edge驱动路径
edge_driver_path = r'D:\Edgedriver\msedgedriver.exe'

### Folder to store screenshots
screenshot_folder = 'screenshots'
os.makedirs(screenshot_folder, exist_ok=True)  # Create folder if it does not exist


# ---------- 读取Excel ----------
df = pd.read_excel(excel_path)
company_names = df.iloc[:, 0].dropna().tolist()

# ---------- 启动Edge ----------
edge_options = Options()
edge_options.add_argument("--start-maximized")
driver = webdriver.Edge(service=Service(edge_driver_path), options=edge_options)

# ---------- 打开PDF ----------
driver.get(f'file:///{pdf_path}')
time.sleep(5)  # 等待PDF加载

# ---------- 创建Word文档 ----------
doc = Document()
doc.add_heading('Sanctions List Search Results', 0)

# ---------- 遍历公司名 ----------
for company in company_names:
    print(f"Searching for: {company}")

    # # 切换到Edge窗口最前
    # pyautogui.hotkey('alt', 'tab')
    # time.sleep(10)

    # 按Ctrl+F调出搜索框
    pyautogui.hotkey('ctrl', 'f')
    time.sleep(2)
    
    # 全选并清空
    pyautogui.hotkey('ctrl', 'a')
    time.sleep(0.5)
    pyautogui.press('backspace')
    time.sleep(0.5)

    # 把公司名复制到剪贴板
    pyperclip.copy(company)
    time.sleep(0.5)

    # 粘贴公司名
    pyautogui.hotkey('ctrl', 'v')
    time.sleep(2)   ############# 根据你的电脑速度可缩短或加长

    # 按Enter开始搜索
    pyautogui.press('enter')

    # 等待Edge完成搜索（视文件大小可适当调大）
    time.sleep(105)    #############待调整

    # 截图
    img = ImageGrab.grab(bbox=screenshot_box)
    screenshot_filename = os.path.join(screenshot_folder, f"{company}.png")
    img.save(screenshot_filename)

    # 添加到Word
    doc.add_heading(company, level=2)
    doc.add_picture(screenshot_filename, width=Inches(5))

# 保存Word
doc.save(word_path)

print("Search and documentation complete!")

driver.quit()
